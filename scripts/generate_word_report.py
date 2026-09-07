"""
generate_word_report.py: Compiles the complete, publication-grade academic project report
into a beautifully formatted Microsoft Word (.docx) document with embedded high-resolution
figures, structured comparison tables, callout boxes, and clinical disclaimers.
"""

import sys
from pathlib import Path

import docx
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


def set_cell_shading(cell, color_hex: str):
    """Applies background color to a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def create_callout_box(doc, text: str, title: str = "IMPORTANT MEDICAL NOTICE", bg_color: str = "FFF1F2", border_color: str = "F43F5E"):
    """Creates a stylized callout alert box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    # Set left border thick, others none
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"⚠️ {title}\n")
    run_title.bold = True
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = RGBColor(185, 28, 28)

    run_body = p.add_run(text)
    run_body.font.size = Pt(9.5)
    run_body.font.color.rgb = RGBColor(51, 65, 85)


def add_styled_heading(doc, text: str, level: int):
    """Adds a heading with consistent academic typography."""
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    run = h.runs[0]
    if level == 1:
        run.font.name = "Calibri"
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 118, 110)  # Teal
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.name = "Calibri"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59)  # Slate
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(71, 85, 105)
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
    return h


def build_academic_word_report(output_path: str = "Explainable_AI_Diabetes_Prediction_Project_Report.docx"):
    doc = Document()

    # Configure Margins (0.8 inches all around)
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # Base typography
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(51, 65, 85)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------------------
    # COVER / HEADER SECTION
    # -------------------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("Explainable AI Diabetes Prediction & Clinical Decision Support System")
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run(
        "A Probability-Calibrated Machine Learning Framework with Dual Local Explainability (SHAP + LIME) "
        "and Personalized Clinical Educational Guidance"
    )
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Metadata Grid Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    col_widths = [Inches(2.5), Inches(4.3)]
    meta_data = [
        ("Project Type", "Final-Year Academic Capstone / Research Engineering Project"),
        ("Clinical Focus", "Type 2 Diabetes Screening & Risk Stratification (PIMA Indians Cohort)"),
        ("System Architecture", "Calibrated XGBoost, TreeExplainer SHAP, LIME Surrogate, Streamlit, ReportLab, SMTP, Gemini"),
        ("Verification Status", "Empirically Validated & Calibrated | 45 Automated Tests Passing (100%)"),
    ]

    for row_idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[row_idx]
        row.cells[0].width = col_widths[0]
        row.cells[1].width = col_widths[1]

        set_cell_shading(row.cells[0], "F1F5F9")
        set_cell_shading(row.cells[1], "F8FAFC")
        set_cell_margins(row.cells[0], top=80, bottom=80, left=120, right=120)
        set_cell_margins(row.cells[1], top=80, bottom=80, left=120, right=120)

        p0 = row.cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.size = Pt(9)
        r0.font.color.rgb = RGBColor(30, 41, 59)

        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v)
        r1.font.size = Pt(9)
        r1.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Callout Box: Medical Disclaimer
    create_callout_box(
        doc,
        "This software and generated report represent an AI-based research prototype for educational and screening decision support. "
        "It is strictly non-diagnostic and does not replace consultation with a qualified medical professional. "
        "Patients must never initiate, alter, or discontinue any medications or treatments without direct physician supervision.",
        title="MANDATORY CLINICAL NON-DIAGNOSTIC NOTICE"
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------------------
    # 1. EXECUTIVE SUMMARY
    # -------------------------------------------------------------------------
    add_styled_heading(doc, "1. Executive Summary", level=1)
    p = doc.add_paragraph(
        "Machine learning models in healthcare often suffer from the classic 'black-box' dilemma: while high-capacity "
        "ensemble tree architectures such as XGBoost achieve strong discrimination, their opaque decision topologies hinder "
        "physician adoption, clinical verification, and patient understanding. This project resolves this dilemma by constructing "
        "a fully transparent, probability-calibrated, and ethically constrained Explainable AI Diabetes Prediction System."
    )
    p = doc.add_paragraph(
        "Building upon the benchmark PIMA Indians Diabetes Database (N=768, 8 numerical features), the engineering core was upgraded "
        "to optimize recall on the diabetic class while ensuring honest, leak-free evaluation. The baseline XGBoost model (holdout ROC-AUC 0.8198, "
        "Accuracy 74.03%) was preserved intact, while a calibrated decision model was engineered that improves sensitivity from 55.56% to 81.48%, "
        "slashing missed diabetic cases (False Negatives) by more than half (from 24 down to 10)."
    )
    p = doc.add_paragraph(
        "To deliver actionable transparency, the system provides dual local explainability: game-theoretic local Shapley attributions "
        "(TreeExplainer) and local linear surrogate approximations (LIME), validated by quantitative agreement metrics (Spearman rank "
        "correlation rho = 0.8810, Top-3 feature overlap 100%). Furthermore, a deterministic clinical recommendation engine translates "
        "the patient's positive local SHAP risk drivers into personalized educational guidance across a multi-channel delivery layer comprising "
        "an interactive 7-tab Streamlit application, publication-grade 9-section ReportLab PDFs, and secure TLS SMTP email dispatch."
    )

    # -------------------------------------------------------------------------
    # 2. CLINICAL MOTIVATION & PROBLEM STATEMENT
    # -------------------------------------------------------------------------
    add_styled_heading(doc, "2. Clinical Motivation & Problem Statement", level=1)
    doc.add_paragraph(
        "Type 2 diabetes mellitus is a progressive metabolic disorder characterized by hyperglycemia resulting from insulin resistance "
        "and relative insulin deficiency. The asymptomatic nature of early-stage dysglycemia frequently leads to delayed diagnosis, at which point "
        "irreversible microvascular and macrovascular damage may have already commenced."
    )
    doc.add_paragraph(
        "Traditional risk assessment questionnaires often lack mathematical personalization, whereas automated black-box classifiers "
        "output opaque probabilities that clinicians cannot verify. Furthermore, generic lifestyle advice fails to engage patients because "
        "it does not illuminate which specific physiological factors contributed to the patient's individual risk tier."
    )
    doc.add_paragraph(
        "This project was undertaken to meet three core research objectives:\n"
        "1. Maximize Screening Sensitivity: Prioritize calibrated recall on diabetic cases to minimize harmful False Negatives.\n"
        "2. Deliver Dual Local Transparency: Deconstruct individual predictions using both exact tree Shapley attributions and local surrogate approximations.\n"
        "3. Bridge Prediction to Actionable Guidance: Synthesize personalized educational wellness recommendations grounded in the specific positive SHAP contributors."
    )

    # -------------------------------------------------------------------------
    # 3. END-TO-END SYSTEM ARCHITECTURE
    # -------------------------------------------------------------------------
    add_styled_heading(doc, "3. End-to-End System Architecture", level=1)
    doc.add_paragraph(
        "The system follows a strictly modular, decoupled 5-tier architecture that guarantees data hygiene, zero leakage, "
        "and high availability:"
    )

    arch_items = [
        ("Tier 1 — Demographic & Clinical Ingestion (UI)", "Collects Patient Name, Biological Sex, Age, and 7 physiological measurements in Streamlit. Dynamic UI logic automatically sets Pregnancies to 0 and disables the input when Male is selected."),
        ("Tier 2 — Leak-Free Preprocessing Core", "Treats implausible zeros in Glucose, Blood Pressure, Skin Thickness, Insulin, and BMI as missing values via MissingZeroHandler. Transforms inputs using median imputer and z-score standard scaler fitted strictly on the training set."),
        ("Tier 3 — Calibrated Machine Learning Engine", "Executes serialized XGBoost inference to generate discrete classifications and well-calibrated class-1 probabilities."),
        ("Tier 4 — Dual Explainability & Recommendation Core", "Computes local signed SHAP attributions via TreeExplainer and fits a LIME tabular surrogate. The deterministic recommendation engine identifies positive upward drivers, applies modifiable vs. non-modifiable rules, and builds personalized guidance."),
        ("Tier 5 — Multi-Channel Delivery Layer", "Renders interactive results across 7 Streamlit tabs, compiles 9-section ReportLab clinical PDFs, dispatches TLS SMTP emails, and optionally queries Google Gemini (gemini-flash-lite-latest) with 100% deterministic offline fallback.")
    ]
    for name, desc in arch_items:
        p = doc.add_paragraph(style="List Bullet")
        r1 = p.add_run(f"{name}: ")
        r1.bold = True
        p.add_run(desc)

    # -------------------------------------------------------------------------
    # 4. DATASET & PREPROCESSING PIPELINE
    # -------------------------------------------------------------------------
    add_styled_heading(doc, "4. Dataset Analysis & Preprocessing Methodology", level=1)
    doc.add_paragraph(
        "The project is evaluated on the benchmark PIMA Indians Diabetes Database (National Institute of Diabetes and Digestive "
        "and Kidney Diseases). The dataset contains 768 patient records with 8 numerical clinical features and a binary Outcome target."
    )

    # Dataset Table
    add_styled_heading(doc, "Table 1: Clinical Feature Specifications & Missing-Zero Rationale", level=2)
    feat_table = doc.add_table(rows=9, cols=4)
    feat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(1.8), Inches(1.2), Inches(1.8), Inches(2.0)]

    headers = ["Feature Name", "Data Type / Unit", "Physiological Normal Range", "Zero Value Handling"]
    for i, h in enumerate(headers):
        cell = feat_table.rows[0].cells[i]
        cell.width = widths[i]
        set_cell_shading(cell, "0F766E")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    feat_rows = [
        ("Pregnancies", "Integer (count)", "0 – 17", "Preserved as biologically valid (0 = nulliparous)"),
        ("Glucose", "Integer (mg/dL)", "70 – 140 mg/dL (fasting/OGTT)", "0 is biologically implausible -> NaN -> Median Imputed"),
        ("BloodPressure", "Integer (mmHg)", "60 – 90 mmHg (diastolic)", "0 is biologically implausible -> NaN -> Median Imputed"),
        ("SkinThickness", "Integer (mm)", "10 – 50 mm (triceps skinfold)", "0 is biologically implausible -> NaN -> Median Imputed"),
        ("Insulin", "Integer (muU/mL)", "16 – 166 muU/mL (2-hour post)", "0 is biologically implausible -> NaN -> Median Imputed"),
        ("BMI", "Float (kg/m²)", "18.5 – 24.9 kg/m²", "0 is biologically implausible -> NaN -> Median Imputed"),
        ("DiabetesPedigreeFunction", "Float (ratio)", "0.08 – 2.42 (genetic score)", "Preserved as calculated (strictly positive)"),
        ("Age", "Integer (years)", "21 – 81 years", "Preserved as recorded (minimum age 21)"),
    ]

    for row_idx, data in enumerate(feat_rows, start=1):
        row = feat_table.rows[row_idx]
        bg = "FFFFFF" if row_idx % 2 != 0 else "F8FAFC"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            cell.width = widths[col_idx]
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(8)
            if col_idx == 0:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------------------
    # 5. MACHINE LEARNING MODELING & BENCHMARKING
    # -------------------------------------------------------------------------
    add_styled_heading(doc, "5. Machine Learning Modeling & Benchmarking", level=1)
    doc.add_paragraph(
        "To establish rigorous baseline comparisons, five algorithms were benchmarked using identical leak-free 5-fold Stratified "
        "Cross-Validation on the training cohort (N=614):"
    )

    # Benchmark Table
    add_styled_heading(doc, "Table 2: Candidate Models 5-Fold Stratified Cross-Validation Benchmark", level=2)
    bench_tbl = doc.add_table(rows=6, cols=5)
    bench_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    b_widths = [Inches(2.2), Inches(1.2), Inches(1.2), Inches(1.1), Inches(1.1)]

    b_headers = ["Algorithm Evaluated", "Mean CV ROC-AUC", "Mean CV Accuracy", "Mean CV Recall", "Mean CV F1"]
    for i, h in enumerate(b_headers):
        cell = bench_tbl.rows[0].cells[i]
        cell.width = b_widths[i]
        set_cell_shading(cell, "0F766E")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    b_data = [
        ("XGBoost (Baseline)", "0.8322 (+/- 0.0237)", "75.89%", "60.29%", "63.81%"),
        ("XGBoost (Hyperparameter Tuned)", "0.8430 (+/- 0.0195)", "76.87%", "67.28%", "67.42%"),
        ("LightGBM Classifier", "0.8315 (+/- 0.0210)", "75.40%", "62.15%", "64.10%"),
        ("Random Forest Classifier", "0.8290 (+/- 0.0245)", "75.24%", "58.88%", "62.40%"),
        ("Logistic Regression", "0.8265 (+/- 0.0220)", "76.22%", "56.54%", "62.50%"),
    ]

    for r_idx, data in enumerate(b_data, start=1):
        row = bench_tbl.rows[r_idx]
        bg = "FFFFFF" if r_idx % 2 != 0 else "F8FAFC"
        for c_idx, text in enumerate(data):
            cell = row.cells[c_idx]
            cell.width = b_widths[c_idx]
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(8)
            if c_idx == 0:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Baseline vs Calibrated Table
    add_styled_heading(doc, "Table 3: Baseline vs. Optimized Holdout Test Performance (N=154)", level=2)
    hold_tbl = doc.add_table(rows=10, cols=4)
    hold_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_widths = [Inches(2.5), Inches(1.4), Inches(1.4), Inches(1.5)]

    h_headers = ["Metric / Clinical Criterion", "Baseline XGBoost", "Optimized XGBoost", "Calibrated XGBoost (tau=0.30)"]
    for i, h in enumerate(h_headers):
        cell = hold_tbl.rows[0].cells[i]
        cell.width = h_widths[i]
        set_cell_shading(cell, "0F766E")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    h_data = [
        ("Accuracy", "74.03%", "73.38%", "75.32% (Highest overall)"),
        ("Recall (Sensitivity on Diabetics)", "55.56% (30 / 54)", "83.33% (45 / 54)", "81.48% (44 / 54)"),
        ("Precision", "65.22%", "58.44%", "61.11%"),
        ("F1-Score", "60.00%", "68.70%", "69.84%"),
        ("F2-Score (Recall Prioritized)", "57.25%", "76.79%", "76.39%"),
        ("ROC-AUC Score", "0.8198", "0.8243", "0.8310"),
        ("PR-AUC Score", "0.6751", "0.6918", "0.7009"),
        ("Brier Score (Probability Calibration)", "0.1812", "0.1745", "0.1670 (Lowest calibration error)"),
        ("False Negatives (Missed Diabetic Cases)", "24 missed patients", "9 missed (-62.5%)", "10 missed (-58.3%)"),
    ]

    for r_idx, data in enumerate(h_data, start=1):
        row = hold_tbl.rows[r_idx]
        bg = "FFFFFF" if r_idx % 2 != 0 else "F8FAFC"
        for c_idx, text in enumerate(data):
            cell = row.cells[c_idx]
            cell.width = h_widths[c_idx]
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(8)
            if c_idx == 0:
                r.bold = True
            if "missed (-" in text or "81.48%" in text or "83.33%" in text:
                r.bold = True
                r.font.color.rgb = RGBColor(15, 118, 110)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Embed Diagnostic Figures
    fig_dir = Path("results/figures")
    if (fig_dir / "roc_curves.png").exists():
        add_styled_heading(doc, "Diagnostic Figures: ROC Curves & Model Discrimination", level=2)
        doc.add_picture(str(fig_dir / "roc_curves.png"), width=Inches(5.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap = doc.add_paragraph("Figure 1: Receiver Operating Characteristic (ROC) curve comparison across candidate models.")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.style.font.size = Pt(8)
        p_cap.style.font.italic = True
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------------------
    # 6. DUAL LOCAL EXPLAINABILITY (SHAP vs. LIME)
    # -------------------------------------------------------------------------
    add_styled_heading(doc, "6. Dual Local Explainability (SHAP vs. LIME)", level=1)
    doc.add_paragraph(
        "To eliminate reliance on a single explainability method, the framework combines game-theoretic local Shapley attributions "
        "(TreeExplainer) with local linear surrogate approximations (LIME). This dual approach validates whether independent interpretability "
        "formulations reach consensus on which physiological variables govern the prediction."
    )

    if (fig_dir / "shap_summary.png").exists():
        doc.add_picture(str(fig_dir / "shap_summary.png"), width=Inches(5.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap = doc.add_paragraph("Figure 2: Global SHAP summary beeswarm plot illustrating feature value distributions and risk impacts.")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.style.font.size = Pt(8)
        p_cap.style.font.italic = True
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    if (fig_dir / "shap_lime_comparison.png").exists():
        doc.add_picture(str(fig_dir / "shap_lime_comparison.png"), width=Inches(5.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap = doc.add_paragraph("Figure 3: Concordance comparison between SHAP Shapley attributions and LIME surrogate weights.")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.style.font.size = Pt(8)
        p_cap.style.font.italic = True
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Quantitative concordance metrics evaluated on holdout predictions demonstrated high consistency:\n"
        "• Spearman Rank Correlation: rho = 0.8810 (p = 0.0039), confirming strong monotonic feature alignment.\n"
        "• Top-3 Feature Overlap (Jaccard Index): 1.00 (100% agreement on the three primary drivers).\n"
        "• Sign Concordance: 100%, indicating perfect agreement on directional risk elevation vs. mitigation."
    )

    # -------------------------------------------------------------------------
    # 7. PERSONALIZED SHAP-BASED RECOMMENDATION ENGINE
    # -------------------------------------------------------------------------
    add_styled_heading(doc, "7. Personalized SHAP-Based Recommendation Engine", level=1)
    doc.add_paragraph(
        "A central innovation of this project is replacing static, generic advice with a deterministic educational guidance engine "
        "(services/recommendations.py). The engine dynamically interrogates the patient's local signed SHAP attributions to generate "
        "tailored wellness recommendations directly connected to the specific variables pushing the model toward elevated risk."
    )

    add_styled_heading(doc, "Table 4: Feature Modifiability & Personalized Clinical Guidance Rules", level=2)
    rec_tbl = doc.add_table(rows=9, cols=3)
    rec_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    r_widths = [Inches(1.8), Inches(1.4), Inches(3.6)]

    rec_headers = ["Clinical Parameter", "Factor Modifiability", "Model Context & Personalized Guidance Rules"]
    for i, h in enumerate(rec_headers):
        cell = rec_tbl.rows[0].cells[i]
        cell.width = r_widths[i]
        set_cell_shading(cell, "0F766E")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    rec_rows = [
        ("Glucose", "Modifiable", "Suggests physician discussion regarding fasting plasma glucose and HbA1c screening; emphasizes dietary fiber, whole grains, and limiting refined sugars. Strictly forbids independent medication changes."),
        ("BMI", "Modifiable", "Encourages sustainable healthy eating patterns and physical activity tailored to fitness. Avoids crash diets and never promises that losing X kg will drop risk by Y%."),
        ("BloodPressure", "Modifiable", "Encourages resting blood pressure monitoring and sodium reduction. Clarifies that a single input reading does not constitute a diagnosis of hypertension."),
        ("Insulin", "Modifiable", "Recommends comprehensive metabolic evaluation with a physician; reminds patient that model uses insulin as a predictive feature, not proof of insulin resistance."),
        ("Age", "NON-MODIFIABLE", "Designates age as non-modifiable; emphasizes healthy aging, age-appropriate routine screenings, and low-impact exercise. Zero advice to 'reduce age'."),
        ("Pregnancies", "NON-MODIFIABLE", "Historical obstetric marker; focuses on general metabolic wellness and maternal health history awareness."),
        ("DiabetesPedigreeFunction", "NON-MODIFIABLE", "Genetic predisposition score; encourages family medical history awareness and proactive lifestyle resilience."),
        ("SkinThickness", "Modifiable Proxy", "Non-diagnostic anthropometric proxy; recommends well-rounded nutrition and general physical activity."),
    ]

    for r_idx, data in enumerate(rec_rows, start=1):
        row = rec_tbl.rows[r_idx]
        bg = "FFFFFF" if r_idx % 2 != 0 else "F8FAFC"
        for c_idx, text in enumerate(data):
            cell = row.cells[c_idx]
            cell.width = r_widths[c_idx]
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(8)
            if c_idx == 0:
                r.bold = True
            if "NON-MODIFIABLE" in text:
                r.bold = True
                r.font.color.rgb = RGBColor(185, 28, 28)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------------------
    # 8. DEMOGRAPHIC METADATA ISOLATION & GENDER-CONDITIONAL UI
    # -------------------------------------------------------------------------
    add_styled_heading(doc, "8. Demographic Metadata Isolation & Gender-Conditional UI", level=1)
    doc.add_paragraph(
        "Because the PIMA Indians Diabetes Dataset contains solely female subjects, introducing demographic fields (Patient Name, "
        "Age, Sex) presents a major architectural challenge: passing demographic fields into the model would alter the mathematical "
        "feature matrix, break serialized model weights, and produce invalid SHAP dimensions."
    )
    doc.add_paragraph(
        "To resolve this, the system strictly isolates demographics:\n"
        "1. Schema Decoupling: FEATURE_ORDER remains strictly locked to the 8 numerical features. Neither Name nor Sex is ever passed to the XGBoost model, scaler, or imputer.\n"
        "2. No Fabricated Attributions: No SHAP value is ever calculated for Name or Sex.\n"
        "3. Gender-Conditional UI Form: In Tab 1 of app.py, selecting Male automatically locks Pregnancies to 0 and disables the input widget with a helpful caption. Selecting Female or Not specified unlocks the field for input."
    )

    # -------------------------------------------------------------------------
    # 9. MULTI-CHANNEL REPORTING & SERVICE LAYER
    # -------------------------------------------------------------------------
    add_styled_heading(doc, "9. Multi-Channel Reporting & Service Layer", level=1)
    doc.add_paragraph(
        "A single shared recommendation function (generate_personalized_recommendations) feeds into all four delivery channels:"
    )
    services_list = [
        ("7-Tab Streamlit Web Application (app.py)", "Provides interactive inputs, real-time probability meter, top-3 SHAP cards, sensitivity sweeps, perturbation tests, and interactive personalized guidance expanders."),
        ("ReportLab Clinical PDF Generator (services/pdf_report.py)", "Compiles an academic-grade 9-section PDF report featuring Patient Information, Clinical Measurements, Risk Assessment, Local SHAP table with explicit directional impact wording, Personalized Summary, Personalized Educational Guidance, and disclaimers."),
        ("Secure SMTP Email Dispatcher (services/email_service.py)", "Enforces RFC 5322 syntax, authenticates via TLS/STARTTLS, and transmits plain-text summaries with the complete PDF attached. Never fakes delivery."),
        ("Grounded LLM Synthesis (services/ai_report.py)", "Connects to Google Gemini (gemini-flash-lite-latest) with strict grounding prompts. If unconfigured or offline, seamlessly falls back to the deterministic clinical template without crashing.")
    ]
    for s_name, s_desc in services_list:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(f"{s_name}: ")
        r.bold = True
        p.add_run(s_desc)

    # -------------------------------------------------------------------------
    # 10. EMPIRICAL ROBUSTNESS & AUTOMATED TESTING
    # -------------------------------------------------------------------------
    add_styled_heading(doc, "10. Empirical Robustness & Automated Test Suite", level=1)
    doc.add_paragraph(
        "To ensure deployment stability, the system was subjected to rigorous perturbation testing and comprehensive unit testing:"
    )
    doc.add_paragraph(
        "• Local Noise Perturbation Testing: Under 50 iterations of +/- 2% Gaussian noise applied across all features, the model demonstrated a 0.00% prediction flip rate, a probability standard deviation of sigma = 0.0084, and a normalized Robustness Index of 1.000 / 1.000.\n"
        "• Automated Test Suite: All 45 automated unit and integration tests pass with 100% success across 7 test modules:"
    )

    test_modules = [
        ("tests/test_personalized_guidance.py", "15 Tests", "Validates personalized guidance, PDF text streams, demographic isolation, and non-diagnostic guardrails."),
        ("tests/test_email.py", "6 Tests", "Validates RFC 5322 syntax, email formatting, and SMTP connections."),
        ("tests/test_reports.py", "4 Tests", "Validates LLM fallback and PDF document compilation."),
        ("tests/test_explainability.py", "6 Tests", "Validates SHAP, LIME, and agreement calculations."),
        ("tests/test_preprocessing.py", "5 Tests", "Validates zero-handling, median imputation, and standard scaling."),
        ("tests/test_models.py", "5 Tests", "Validates baseline holdout metrics and artifact loading."),
        ("tests/test_validation.py", "4 Tests", "Validates boundary enforcement and physiological range rules."),
    ]

    for t_file, t_count, t_desc in test_modules:
        p = doc.add_paragraph(style="List Bullet")
        r1 = p.add_run(f"{t_file} ({t_count}): ")
        r1.bold = True
        p.add_run(t_desc)

    # -------------------------------------------------------------------------
    # 11. ETHICAL GOVERNANCE & LIMITATIONS
    # -------------------------------------------------------------------------
    add_styled_heading(doc, "11. Ethical Considerations & Project Limitations", level=1)
    doc.add_paragraph(
        "1. Population Scope: Training data reflects female Pima Indian patients >= 21 years. Caution is advised when generalizing predictions to diverse demographics.\n"
        "2. Imputation Variance: Zeros in metabolic features are imputed with cohort medians, which may attenuate extreme physiological variance.\n"
        "3. Non-Causality of SHAP: SHAP attributions reflect mathematical reliance in the tree ensemble, not direct biological causation.\n"
        "4. Non-Diagnostic Boundary: The system is an academic research prototype and not an FDA-cleared diagnostic device."
    )

    # -------------------------------------------------------------------------
    # 12. CONCLUSION & FUTURE WORK
    # -------------------------------------------------------------------------
    add_styled_heading(doc, "12. Conclusion & Future Roadmap", level=1)
    doc.add_paragraph(
        "This project successfully transforms a baseline XGBoost classifier into an academic-grade, publication-ready Explainable AI "
        "decision-support system. By uniting calibrated machine learning, dual local explainability (SHAP + LIME), quantitative agreement "
        "metrics, personalized educational guidance, and automated multi-channel reporting, the system establishes a comprehensive model "
        "for transparent clinical machine learning."
    )
    doc.add_paragraph(
        "Future enhancements include expanding to multi-cohort datasets (e.g., NHANES, UK Biobank), incorporating longitudinal sequential "
        "monitoring, and exporting outputs as HL7 FHIR clinical observation resources for electronic health record (EHR) integration."
    )

    # Save document
    out_p = Path(output_path)
    out_p.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_p))
    print(f"Word document created successfully at: {out_p.resolve()} (Size: {out_p.stat().st_size} bytes)")
    return str(out_p.resolve())


if __name__ == "__main__":
    out_file = sys.argv[1] if len(sys.argv) > 1 else "Explainable_AI_Diabetes_Prediction_Project_Report.docx"
    build_academic_word_report(out_file)
